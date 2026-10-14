"""
apps.accelerator.intake
=======================
Reading a program's own application data into Pitchmic: the spreadsheet and
JSON readers, the column mapping, and the per-row verdicts a dry-run preview
shows before anything is written.

Two rules govern everything here.

**Nothing is guessed on a founder's behalf.** A header is mapped onto a field
only when it folds onto exactly one known field after case, spacing and
punctuation are removed. Anything else is carried into the payload untouched and
named, so a column is never dropped silently or filed under a heading that
happens to look similar. The same applies to values: an empty cell stays empty
rather than becoming a plausible default.

**A row is described, not judged.** ``issues`` lists what could not be resolved
on the way in — a missing email, an address that is not an address, the same
company twice in one file. It is a checklist for the person deciding, and it
decides nothing: an issue never rejects a row, and a clean row is still held
until a human accepts it.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import secrets
from typing import Any, Iterable

from django.core.validators import EmailValidator

# The four fields a staged row understands structurally. Everything a program's
# form sends beyond these is kept in the payload, because the application spine
# only resolves identity and company; the rest is a human's reading material.
CANONICAL_FIELDS: tuple[str, ...] = ("founder_email", "founder_name", "company_name", "external_id")

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "founder_email": ("email", "e_mail", "email_address", "founder_email", "applicant_email", "contact_email", "primary_email"),
    "founder_name": ("name", "full_name", "founder_name", "applicant_name", "contact_name", "primary_founder", "founder"),
    "company_name": ("company", "company_name", "startup", "startup_name", "business", "business_name", "organization", "organisation", "venture", "venture_name", "trading_name"),
    "external_id": ("id", "ref", "reference", "reference_id", "external_id", "submission_id", "application_id", "row_id", "record_id"),
}

FIELD_LABELS: dict[str, str] = {
    "founder_email": "Founder email",
    "founder_name": "Founder name",
    "company_name": "Company name",
    "external_id": "Their reference",
}

MAX_ROWS = 2000
MAX_CELL_CHARS = 4000
MAX_PAYLOAD_BYTES = 64 * 1024

# Issues a person should read but which do not stand between a row and being
# filed: a second copy inside one file, and a repeat of something the queue
# already holds. Neither says the founder is wrong, so neither blocks acceptance.
NON_BLOCKING_CODES = ("duplicate_in_file", "already_in_queue")

_email_validator = EmailValidator()


class IntakeFormatError(Exception):
    """The file cannot be read as a table at all — a format problem, not a data problem."""


def _fold(header: Any) -> str:
    """Case, spacing and punctuation removed, for matching a header to a field."""
    text = str(header or "").strip().casefold()
    text = re.sub(r"[^a-z0-9]+", "_", text, flags=re.ASCII)
    return text.strip("_")


def canonical_for(header: Any) -> str | None:
    """The one canonical field this header names, or None.

    None on a near miss is deliberate: "email_verified" and "company_website"
    resemble known headers but mean something else, and filing an address under
    the wrong heading is a corruption nobody notices later.
    """
    folded = _fold(header)
    if not folded:
        return None
    for field, aliases in FIELD_ALIASES.items():
        if folded == _fold(field) or folded in aliases:
            return field
    return None


def suggest_mapping(headers: Iterable[Any]) -> dict[str, str]:
    """{canonical field: the header that supplies it} for a first draft of the mapping."""
    by_folded: dict[str, Any] = {}
    for header in headers:
        folded = _fold(header)
        if folded and folded not in by_folded:
            by_folded[folded] = header
    suggested = {}
    for field, aliases in FIELD_ALIASES.items():
        candidates = [_fold(field)] + list(aliases)
        for candidate in candidates:
            if candidate in by_folded:
                suggested[field] = by_folded[candidate]
                break
    return suggested


def _cell_text(value: Any) -> str:
    """A cell as text. A whole number from a spreadsheet arrives as a float."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        # Only JSON imports produce these; a CSV or XLSX cell cannot hold a
        # nested object, so it is kept readable rather than silently dropped.
        return json.dumps(value, default=str)[:MAX_CELL_CHARS]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:MAX_CELL_CHARS]


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_table_upload(django_file) -> tuple[list[str], list[list[str]]]:
    """(headers, rows) from an uploaded .csv/.tsv/.xlsx/.xlsm/.json file."""
    name = (getattr(django_file, "name", "") or "").lower()
    raw = django_file.read()
    if isinstance(raw, str):
        raw = raw.encode("utf-8")
    if not raw:
        raise IntakeFormatError("That file is empty.")
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(raw)
    if name.endswith((".csv", ".tsv", ".txt")):
        return _read_delimited(_decode(raw), delimiter="\t" if name.endswith(".tsv") else None)
    if name.endswith(".json"):
        return _read_json(_decode(raw))
    raise IntakeFormatError(
        "No reader for that format. Accepted: CSV, TSV, XLSX, XLSM, JSON. "
        "If the data lives in another system, post it to the cycle's intake key instead."
    )


def _read_xlsx(raw: bytes) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover — openpyxl is a declared dependency
        raise IntakeFormatError("Excel reading needs the openpyxl package, which is not installed.")
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises its own opaque exceptions
        raise IntakeFormatError(f"That workbook could not be opened ({type(exc).__name__}).")
    try:
        for sheet in book.worksheets:
            grid = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            grid = [row for row in grid if any(cell for cell in row)]
            if grid:
                return _finalize(grid)
        raise IntakeFormatError("That workbook has no rows to read.")
    finally:
        book.close()


def _read_delimited(text: str, delimiter: str | None = None) -> tuple[list[str], list[list[str]]]:
    if delimiter == "\t":
        dialect = csv.excel_tab
    elif delimiter is not None:
        dialect = csv.excel
    else:
        try:
            dialect = csv.Sniffer().sniff(text[:4096])
        except csv.Error:
            dialect = csv.excel
    grid = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    grid = [row for row in grid if any(cell for cell in row)]
    if not grid:
        raise IntakeFormatError("That file has no rows to read.")
    return _finalize(grid)


def _read_json(text: str) -> tuple[list[str], list[list[str]]]:
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise IntakeFormatError(f"That file is not valid JSON ({exc.msg} at character {exc.pos}).")
    if isinstance(data, dict):
        data = data.get("rows") or data.get("submissions") or data.get("data")
    if not isinstance(data, list):
        raise IntakeFormatError("Expected a JSON array of submission objects.")
    objects = [item if isinstance(item, dict) else {"value": item} for item in data]
    headers: list[str] = []
    for item in objects:
        for key in item:
            if key not in headers:
                headers.append(str(key))
    if not headers:
        raise IntakeFormatError("That JSON has no submission objects in it.")
    grid = [headers] + [[_cell_text(item.get(key)) for key in headers] for item in objects]
    return _finalize(grid)


def _finalize(grid: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    headers, rows = grid[0], grid[1:]
    named = [header if header.strip() else f"column_{index + 1}" for index, header in enumerate(headers)]
    width = len(named)
    normalised = [(row + [""] * width)[:width] for row in rows]
    return named, normalised[:MAX_ROWS]


def normalize_row(headers: list[str], row: list[str], mapping: dict[str, str]) -> dict[str, Any]:
    """One row of a table as a staged submission: fields, payload, and its issues.

    ``mapping`` is {canonical field: header}. Columns the mapping does not name
    are carried into the payload under their own heading rather than discarded.
    """
    column_of = {field: _column_index(headers, mapping.get(field)) for field in CANONICAL_FIELDS}

    def take(field: str) -> str:
        index = column_of[field]
        return row[index].strip() if index is not None and index < len(row) else ""

    values = {field: take(field) for field in CANONICAL_FIELDS}
    mapped_headers = {h for h in mapping.values() if h}
    payload = {
        header: (row[index].strip() if index < len(row) else "")
        for index, header in enumerate(headers)
        if header not in mapped_headers and index < len(row) and row[index].strip()
    }

    issues: dict[str, dict[str, str]] = {}
    email = values["founder_email"]
    if email:
        try:
            _email_validator(email)
        except Exception:
            issues["email_invalid"] = {
                "code": "email_invalid", "field": "founder_email",
                "message": f"“{email}” is not an email address, so it cannot name an account.",
            }
        values["founder_email"] = email.lower()

    for field, label in (("founder_email", "founder email"), ("company_name", "company name")):
        if not values[field]:
            issues[f"missing_{field}"] = {
                "code": f"missing_{field}", "field": field,
                "message": f"No {label} on this row — it can be staged, but not accepted.",
            }

    return {
        "founder_email": values["founder_email"],
        "founder_name": values["founder_name"],
        "company_name": values["company_name"],
        "external_id": values["external_id"] or None,
        "payload": payload,
        "issues": list(issues.values()),
    }


def _column_index(headers: list[str], header: str | None) -> int | None:
    if not header:
        return None
    try:
        return headers.index(header)
    except ValueError:
        return None


def normalize_table(headers: list[str], rows: list[list[str]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    """Every row, numbered from the header, with duplicates inside the file marked.

    A repeated (email, company) pair inside one file is almost always the same
    founder sent twice, so the second row says so — but it is still staged, because
    a person may have deliberately listed them again.
    """
    seen: set[tuple[str, str]] = set()
    prepared: list[dict[str, Any]] = []
    for offset, row in enumerate(rows):
        record = normalize_row(headers, row, mapping)
        record["row_number"] = offset + 2  # +1 for the header row, +1 for 1-based display
        key = (record["founder_email"].casefold(), record["company_name"].casefold())
        if key != ("", "") and key in seen:
            record["issues"].append({
                "code": "duplicate_in_file", "field": "company_name",
                "message": "The same founder and company appear twice in this file.",
            })
        if key != ("", ""):
            seen.add(key)
        prepared.append(record)
    return prepared


def identity_of(record: dict[str, Any]) -> tuple[str, str]:
    """The (email, company) pair a queue row is recognised by."""
    return (record["founder_email"].strip().casefold(), record["company_name"].strip().casefold())


def queue_identities(program_id, cycle_id) -> set[tuple[str, str]]:
    """Founder-and-company pairs already waiting in this queue.

    Only staged rows count. A dismissed row is free to arrive again — a person
    correcting a file expects its rows back — and an accepted one is answered at
    accept time by the funnel's own uniqueness, not by the queue. A row with
    neither half of the pair has nothing to be recognised by, so it is not
    listed; a row that arrived with an email but no company still has the email,
    and a re-import of that same row is the same incomplete row.
    """
    from .models_os import ExternalIntakeSubmission

    rows = ExternalIntakeSubmission.objects.filter(
        program_id=program_id, cycle_id=cycle_id, status="staged",
    ).values_list("founder_email", "company_name")
    return {
        (email.strip().casefold(), company.strip().casefold())
        for email, company in rows
        if email.strip() or company.strip()
    }


def mark_queue_repeats(records: list[dict[str, Any]], in_queue: set[tuple[str, str]]) -> int:
    """Flag the rows this queue is already holding. Returns how many were flagged.

    Flagged rather than dropped: a re-run of the same file must not quietly
    double the queue, and it must not quietly lose a row either, so the row is
    reported and left for a person to decide what the repeat means.
    """
    repeats = 0
    for record in records:
        if identity_of(record) in in_queue:
            record["issues"].append({
                "code": "already_in_queue", "field": "company_name",
                "message": "This founder and company are already waiting in the queue.",
            })
            repeats += 1
    return repeats


# ──────────────────────────────────────────────────────────────────────────────
#  The cycle's intake key
# ──────────────────────────────────────────────────────────────────────────────

INTAKE_KEY_PREFIX = "pmit_"


def issue_intake_key() -> tuple[str, str, str]:
    """Return ``(key, hash, hint)`` for a cycle's intake key.

    Only the hash is stored, and the plaintext is handed back exactly once — it
    is what a program pastes into their own form. A leaked database therefore
    does not contain working credentials, and a key that has been mislaid is
    rotated rather than recovered, which is what a revocation should mean.
    """
    key = INTAKE_KEY_PREFIX + secrets.token_urlsafe(24)
    return key, hash_intake_key(key), key[:16]


def hash_intake_key(key: str) -> str:
    return hashlib.sha256((key or "").encode("utf-8")).hexdigest()


def payload_too_large(payload: Any) -> bool:
    """Whether a posted body is too big to keep verbatim.

    A staging row stores the submission as it arrived, so an unlimited body turns
    a queue into a disk. Anything over the budget is refused outright rather
    than truncated: half a founder's answer is worse than no answer.
    """
    try:
        return len(json.dumps(payload, default=str).encode("utf-8")) > MAX_PAYLOAD_BYTES
    except (TypeError, ValueError):
        return True


# ──────────────────────────────────────────────────────────────────────────────
#  Accepting a staged submission
# ──────────────────────────────────────────────────────────────────────────────

class IntakeNotReady(Exception):
    """The row is missing something the schema cannot do without."""

    def __init__(self, gaps: list[str]):
        self.gaps = gaps
        super().__init__("; ".join(gaps) or "This submission is not ready to accept.")


class IntakeAlreadyFiled(Exception):
    """That founder's company is already in this program."""


def accept_submission(submission, *, decided_by):
    """Turn one staged row into a real application. Returns ``(application, founder_created)``.

    The account question is settled here by a named person, which is the whole
    point of staging: an email that arrived from a website form is not evidence
    of an identity, but a director filing it under that email is a decision with
    an author. The account is opened unable to sign in — no password, and
    ``is_verified`` false — so claiming it has to come through the ordinary
    verification path. A founder who already has an account is reused, never
    duplicated: ``FounderApplicationProfile.user`` is a OneToOne, so a second one
    would not be a fresh start but a fork that cannot be merged.

    Existing profile data is never overwritten by imported text. An accepted row
    that disagrees with what the founder typed themselves is worth a conversation,
    not a silent UPDATE.
    """
    from django.db import transaction
    from django.utils import timezone

    from apps.users.models import User
    from .models import (
        AcceleratorApplication,
        FounderApplicationProfile,
        StartupApplicationProfile,
    )

    gaps = submission.blocking_gaps()
    if gaps:
        raise IntakeNotReady(gaps)
    if submission.status == "accepted" and submission.application_id:
        raise IntakeAlreadyFiled(
            f"Already filed as application #{submission.application_id}."
        )

    email = submission.founder_email.strip()
    company = submission.company_name.strip()

    with transaction.atomic():
        founder = User.objects.filter(email__iexact=email).first()
        founder_created = founder is None
        if founder is None:
            first, _, rest = submission.founder_name.strip().partition(" ")
            founder = User.objects.create_user(
                email=email,
                password=None,  # unusable until they claim the account
                first_name=first,
                last_name=rest,
                role=User.ROLE_STARTUP,
                is_verified=False,
            )

        profile, _ = FounderApplicationProfile.objects.get_or_create(
            user=founder,
            defaults={"full_name": submission.founder_name.strip() or founder.email},
        )

        startup = (
            StartupApplicationProfile.objects.filter(
                founder_profile=profile, company_name__iexact=company
            ).first()
        )
        if startup is None:
            startup = StartupApplicationProfile.objects.create(
                founder_profile=profile, company_name=company
            )

        if AcceleratorApplication.objects.filter(
            program=submission.program, startup_profile=startup
        ).exists():
            raise IntakeAlreadyFiled(
                f"{company} is already in this program, so it was not filed twice."
            )

        application = AcceleratorApplication.objects.create(
            program=submission.program,
            startup_profile=startup,
            submitted_by=founder,
            cycle=submission.cycle,
            status="SUBMITTED",
            submitted_at=timezone.now(),
        )

        submission.status = "accepted"
        submission.decided_by = decided_by
        submission.decided_at = timezone.now()
        submission.application = application
        submission.save(
            update_fields=["status", "decided_by", "decided_at", "application", "updated_at"]
        )

    return application, founder_created
